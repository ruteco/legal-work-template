#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
md2docx — выгрузка Markdown в Word (.docx) с оформлением под российские
юридические документы (А4, Times New Roman 14, поля 30/15/20/20 мм,
полуторный интервал, абзацный отступ 12,5 мм).

Использование:
    python3 md2docx.py вход.md выход.docx [опции]

Поддерживаемая разметка — см. README.md в этой же папке.
"""

import argparse
import os
import re
import sys

try:
    from docx import Document
    from docx.shared import Pt, Mm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml.ns import qn
except ImportError:
    sys.exit(
        "Не найден python-docx. Установите: pip3 install --user python-docx"
    )


# --------------------------------------------------------------------------
# Разбор inline-разметки
# --------------------------------------------------------------------------

# **жирный**, *курсив*, `моноширинный`
#
# Подчёркивания (__жирный__ / _курсив_) сознательно НЕ поддерживаются:
# в юридических документах строки для заполнения вида «___» и подписи
# «_________ / И.И. Иванов» встречаются постоянно, и такая разметка
# съедала бы их.
_INLINE_RE = re.compile(
    r"(\*\*.+?\*\*|\*[^*\n]+?\*|`[^`\n]+?`)",
    re.DOTALL,
)


def parse_inline(text):
    """Возвращает список (текст, bold, italic, code)."""
    out = []
    for part in _INLINE_RE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            out.append((part[2:-2], True, False, False))
        elif part.startswith("`") and part.endswith("`"):
            out.append((part[1:-1], False, False, True))
        elif part.startswith("*") and part.endswith("*"):
            out.append((part[1:-1], False, True, False))
        else:
            out.append((part, False, False, False))
    return out


def strip_md_links(text):
    """[подпись](url) -> подпись; голые <url> -> url."""
    text = re.sub(r"\[([^\]]*)\]\([^)]*\)", r"\1", text)
    text = re.sub(r"<((?:https?|file)://[^>]+)>", r"\1", text)
    return text


# --------------------------------------------------------------------------
# Построение документа
# --------------------------------------------------------------------------


class Builder:
    def __init__(self, opts):
        self.o = opts
        self.doc = Document()
        self._setup_page()
        self._setup_styles()

    def _setup_page(self):
        s = self.doc.sections[0]
        s.top_margin = Mm(self.o.margin_top)
        s.bottom_margin = Mm(self.o.margin_bottom)
        s.left_margin = Mm(self.o.margin_left)
        s.right_margin = Mm(self.o.margin_right)
        self.text_width_mm = (
            210 - self.o.margin_left - self.o.margin_right
        )

    def _setup_styles(self):
        st = self.doc.styles["Normal"]
        st.font.name = self.o.font
        st.font.size = Pt(self.o.size)
        rpr = st.element.get_or_add_rPr()
        rf = rpr.get_or_add_rFonts()
        rf.set(qn("w:eastAsia"), self.o.font)
        rf.set(qn("w:cs"), self.o.font)
        pf = st.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.line_spacing = self.o.spacing

    def _run(self, par, text, bold=False, italic=False, code=False, size=None):
        r = par.add_run(text)
        font = self.o.mono_font if code else self.o.font
        r.font.name = font
        r.font.size = Pt(size if size else (self.o.size - 1 if code else self.o.size))
        r.bold = bold
        r.italic = italic
        rpr = r._element.get_or_add_rPr()
        rf = rpr.get_or_add_rFonts()
        rf.set(qn("w:eastAsia"), font)
        rf.set(qn("w:cs"), font)
        return r

    def _par(self, align="just", indent=None, after=None, before=0,
             left_indent=None, hanging=None, spacing=None):
        p = self.doc.add_paragraph()
        p.alignment = {
            "just": WD_ALIGN_PARAGRAPH.JUSTIFY,
            "center": WD_ALIGN_PARAGRAPH.CENTER,
            "left": WD_ALIGN_PARAGRAPH.LEFT,
            "right": WD_ALIGN_PARAGRAPH.RIGHT,
        }[align]
        pf = p.paragraph_format
        if indent:
            pf.first_line_indent = Mm(indent)
        if left_indent is not None:
            pf.left_indent = Mm(left_indent)
        if hanging is not None:
            pf.first_line_indent = Mm(-hanging)
        pf.space_after = Pt(self.o.para_space if after is None else after)
        pf.space_before = Pt(before)
        if spacing is not None:
            pf.line_spacing = spacing
        return p

    def add_runs(self, par, text):
        for t, b, i, c in parse_inline(strip_md_links(text)):
            self._run(par, t, bold=b, italic=i, code=c)

    # -- блочные элементы --------------------------------------------------

    def title(self, text):
        p = self._par(align="center", after=self.o.para_space + 4)
        for t, b, i, c in parse_inline(strip_md_links(text)):
            self._run(p, t, bold=True, italic=i, code=c)

    def heading(self, text, level):
        p = self._par(
            align="just",
            indent=self.o.indent if self.o.indent else None,
            before=10 if level == 2 else 6,
            after=self.o.para_space,
        )
        for t, b, i, c in parse_inline(strip_md_links(text)):
            self._run(p, t, bold=True, italic=i, code=c)

    def para(self, text, align="just", indent=True):
        p = self._par(
            align=align,
            indent=self.o.indent if (indent and self.o.indent) else None,
        )
        self.add_runs(p, text)

    def header_block(self, lines):
        """Шапка процессуального документа — блок в правой части листа."""
        left = self.text_width_mm - self.o.header_width
        for ln in lines:
            if not ln.strip():
                p = self._par(align="left", left_indent=left, after=4, spacing=1.0)
                self._run(p, "")
                continue
            p = self._par(align="left", left_indent=left, after=0, spacing=1.0)
            for t, b, i, c in parse_inline(strip_md_links(ln)):
                self._run(p, t, bold=b, italic=i, code=c,
                          size=self.o.header_size)
        self._par(after=self.o.para_space + 10).add_run("")

    def centered_block(self, lines):
        for ln in lines:
            p = self._par(align="center", after=2)
            self.add_runs(p, ln)

    def right_block(self, lines):
        for ln in lines:
            p = self._par(align="right", after=2)
            self.add_runs(p, ln)

    def list_item(self, marker, text, level=0):
        step = (self.o.indent if self.o.indent else 8) + 5
        p = self._par(
            align="just",
            left_indent=step * (level + 1),
            hanging=step,
        )
        self._run(p, marker + " ")
        self.add_runs(p, text)

    def page_break(self):
        p = self.doc.add_paragraph()
        p.add_run().add_break(WD_BREAK.PAGE)

    def table(self, rows, header=True):
        ncols = max(len(r) for r in rows)
        t = self.doc.add_table(rows=0, cols=ncols)
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        col_mm = self.text_width_mm / ncols
        for ri, row in enumerate(rows):
            cells = t.add_row().cells
            for ci in range(ncols):
                cell = cells[ci]
                cell.width = Mm(col_mm)
                cell.paragraphs[0].clear() if hasattr(
                    cell.paragraphs[0], "clear"
                ) else None
                par = cell.paragraphs[0]
                par.paragraph_format.space_after = Pt(0)
                par.paragraph_format.line_spacing = 1.0
                par.alignment = (
                    WD_ALIGN_PARAGRAPH.CENTER if (header and ri == 0)
                    else WD_ALIGN_PARAGRAPH.LEFT
                )
                txt = row[ci] if ci < len(row) else ""
                for tx, b, i, c in parse_inline(strip_md_links(txt)):
                    self._run(par, tx, bold=(b or (header and ri == 0)),
                              italic=i, code=c, size=self.o.size - 2)

    def save(self, path):
        self.doc.save(path)


# --------------------------------------------------------------------------
# Разбор Markdown
# --------------------------------------------------------------------------

FENCE_RE = re.compile(r"^:::\s*(\w+)?\s*$")
TABLE_SEP_RE = re.compile(r"^\s*\|?[\s:-]*-[-\s:|]*\|?\s*$")
ULI_RE = re.compile(r"^\s*[-*+]\s+(.*)$")
OLI_RE = re.compile(r"^\s*(\d+[.)])\s+(.*)$")
PAGEBREAK_RE = re.compile(r"^\s*(\\newpage|<!--\s*pagebreak\s*-->)\s*$", re.I)


def indent_level(line, per_level=2, max_level=3):
    """Уровень вложенности пункта списка по ведущим пробелам (таб = 4)."""
    n = 0
    for ch in line:
        if ch == " ":
            n += 1
        elif ch == "\t":
            n += 4
        else:
            break
    return min(n // per_level, max_level)


def split_table_row(line):
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [c.strip() for c in line.split("|")]


def convert(md_text, b):
    lines = md_text.replace("\r\n", "\n").split("\n")
    i = 0
    n = len(lines)
    para_buf = []

    def flush():
        if para_buf:
            b.para(" ".join(x.strip() for x in para_buf))
            para_buf.clear()

    while i < n:
        line = lines[i]
        stripped = line.strip()

        # ::: блоки
        m = FENCE_RE.match(stripped)
        if m and m.group(1):
            kind = m.group(1).lower()
            flush()
            block = []
            i += 1
            while i < n and not FENCE_RE.match(lines[i].strip()):
                block.append(lines[i].rstrip())
                i += 1
            i += 1  # закрывающий :::
            if kind in ("header", "шапка"):
                b.header_block(block)
            elif kind in ("center", "центр"):
                b.centered_block([x for x in block if x.strip()])
            elif kind in ("right", "право", "подпись"):
                b.right_block([x for x in block if x.strip()])
            else:  # неизвестный блок — как обычные абзацы
                for x in block:
                    if x.strip():
                        b.para(x)
            continue

        # разрыв страницы
        if PAGEBREAK_RE.match(stripped):
            flush()
            b.page_break()
            i += 1
            continue

        # пустая строка
        if not stripped:
            flush()
            i += 1
            continue

        # горизонтальная линия / разделитель — пропускаем
        if re.match(r"^\s*(\*\s*){3,}$|^\s*-{3,}\s*$|^\s*_{3,}\s*$", line):
            flush()
            i += 1
            continue

        # заголовки
        hm = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if hm:
            flush()
            level = len(hm.group(1))
            if level == 1:
                b.title(hm.group(2))
            else:
                b.heading(hm.group(2), level)
            i += 1
            continue

        # таблица
        if "|" in line and i + 1 < n and TABLE_SEP_RE.match(lines[i + 1]):
            flush()
            rows = [split_table_row(line)]
            i += 2
            while i < n and "|" in lines[i] and lines[i].strip():
                rows.append(split_table_row(lines[i]))
                i += 1
            b.table(rows, header=True)
            continue

        # списки (уровень вложенности — по отступу: 2 пробела = 1 уровень)
        om = OLI_RE.match(line)
        if om:
            flush()
            b.list_item(om.group(1), om.group(2), level=indent_level(line))
            i += 1
            continue
        um = ULI_RE.match(line)
        if um:
            flush()
            lvl = indent_level(line)
            b.list_item("—" if lvl == 0 else "•", um.group(1), level=lvl)
            i += 1
            continue

        # обычный текст
        para_buf.append(line)
        i += 1

    flush()


# --------------------------------------------------------------------------
# CLI
# --------------------------------------------------------------------------


def main():
    ap = argparse.ArgumentParser(
        description="Выгрузка Markdown в Word (.docx) с оформлением "
                    "под российские юридические документы.",
        formatter_class=argparse.RawDescriptionHelpFormatter,
    )
    ap.add_argument("input", help="исходный .md файл")
    ap.add_argument("output", nargs="?", help="результат .docx "
                                             "(по умолчанию — рядом с исходником)")
    ap.add_argument("--font", default="Times New Roman", help="основной шрифт")
    ap.add_argument("--mono-font", default="Courier New",
                    help="шрифт для `моноширинных` фрагментов")
    ap.add_argument("--size", type=float, default=14, help="кегль, пт (по умолч. 14)")
    ap.add_argument("--header-size", type=float, default=11,
                    help="кегль шапки, пт (по умолч. 11)")
    ap.add_argument("--spacing", type=float, default=1.5,
                    help="межстрочный интервал (по умолч. 1.5)")
    ap.add_argument("--indent", type=float, default=12.5,
                    help="абзацный отступ, мм; 0 — без отступа")
    ap.add_argument("--para-space", type=float, default=6,
                    help="интервал после абзаца, пт")
    ap.add_argument("--header-width", type=float, default=93,
                    help="ширина блока шапки, мм")
    ap.add_argument("--margin-left", type=float, default=30)
    ap.add_argument("--margin-right", type=float, default=15)
    ap.add_argument("--margin-top", type=float, default=20)
    ap.add_argument("--margin-bottom", type=float, default=20)
    opts = ap.parse_args()

    if not os.path.isfile(opts.input):
        sys.exit("Файл не найден: %s" % opts.input)

    out = opts.output or os.path.splitext(opts.input)[0] + ".docx"

    with open(opts.input, encoding="utf-8") as f:
        md = f.read()

    b = Builder(opts)
    convert(md, b)
    b.save(out)
    print("Готово: %s" % out)


if __name__ == "__main__":
    main()
